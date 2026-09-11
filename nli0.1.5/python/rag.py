from __future__ import annotations

import csv
import math
import re
import threading
from dataclasses import dataclass
from pathlib import Path

import config


@dataclass(frozen=True)
class RagRecord:
    source: str
    location: str
    fields: dict[str, str]
    text: str
    always_include: bool = False


_cache_lock = threading.Lock()
_cached_signature: tuple | None = None
_cached_records: list[RagRecord] = []


def _is_true(value: object) -> bool:
    return str(value).strip().lower() in {
        "1", "true", "yes", "y", "on", "use", "active", "required"
    }


def _clean_value(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _normalise(text: str) -> str:
    return re.sub(r"\s+", " ", str(text).lower()).strip()


def _tokens(text: str) -> set[str]:
    """Creates English/code tokens together with Korean-character 2-grams."""
    normalised = _normalise(text)
    result = set(re.findall(r"[a-z0-9_\.\-]+|[\uac00-\ud7a3]+", normalised))

    for korean_word in re.findall(r"[\uac00-\ud7a3]+", normalised):
        if len(korean_word) == 1:
            result.add(korean_word)
        else:
            result.update(
                korean_word[index:index + 2]
                for index in range(len(korean_word) - 1)
            )

    return {token for token in result if token}


def _record_from_fields(
    source: str,
    location: str,
    fields: dict[str, object],
) -> RagRecord | None:
    cleaned = {
        str(key).strip(): _clean_value(value)
        for key, value in fields.items()
        if str(key).strip() and _clean_value(value)
    }

    if not cleaned:
        return None

    enabled_value = cleaned.get("enabled", cleaned.get("use", "true"))
    if enabled_value and not _is_true(enabled_value):
        return None

    always_value = cleaned.get(
        "always_include",
        cleaned.get("always_include_localized", "false"),
    )

    text = " | ".join(
        f"{key}={value}"
        for key, value in cleaned.items()
        if key not in {"enabled", "use", "always_include", "always_include_localized"}
    )

    if not text:
        return None

    return RagRecord(
        source=source,
        location=location,
        fields=cleaned,
        text=text,
        always_include=_is_true(always_value),
    )


def _load_csv(path: Path) -> list[RagRecord]:
    records: list[RagRecord] = []

    with path.open("r", encoding="utf-8-sig", newline="") as file:
        reader = csv.DictReader(file)

        for row_number, row in enumerate(reader, start=2):
            record = _record_from_fields(
                source=path.name,
                location=f"row {row_number}",
                fields=row,
            )
            if record is not None:
                records.append(record)

    return records


def _load_xlsx(path: Path) -> list[RagRecord]:
    try:
        from openpyxl import load_workbook
    except ImportError as error:
        raise RuntimeError(
            "openpyxl in requirements.txt is required to use XLSX RAG."
        ) from error

    records: list[RagRecord] = []
    workbook = load_workbook(path, read_only=True, data_only=True)

    try:
        for worksheet in workbook.worksheets:
            rows = worksheet.iter_rows(values_only=True)
            header_row = next(rows, None)

            if header_row is None:
                continue

            headers = [_clean_value(value) for value in header_row]

            for row_number, values in enumerate(rows, start=2):
                fields = {
                    headers[index]: values[index]
                    for index in range(min(len(headers), len(values)))
                    if headers[index]
                }

                record = _record_from_fields(
                    source=path.name,
                    location=f"{worksheet.title}!row {row_number}",
                    fields=fields,
                )
                if record is not None:
                    records.append(record)

    finally:
        workbook.close()

    return records


def _load_txt(path: Path) -> list[RagRecord]:
    text = path.read_text(encoding="utf-8", errors="replace")
    paragraphs = [
        paragraph.strip()
        for paragraph in re.split(r"\n\s*\n", text)
        if paragraph.strip()
    ]

    return [
        RagRecord(
            source=path.name,
            location=f"paragraph {index}",
            fields={"content": paragraph},
            text=paragraph,
        )
        for index, paragraph in enumerate(paragraphs, start=1)
    ]


def _source_paths() -> list[Path]:
    paths: list[Path] = []

    for filename in config.RAG_SOURCE_FILES:
        path = config.RAG_DIRECTORY / filename
        if path.is_file():
            paths.append(path)
        else:
            print(f"[RAG WARNING] Data file not found: {path}")

    return paths


def _signature(paths: list[Path]) -> tuple:
    return tuple(
        (str(path), path.stat().st_mtime_ns, path.stat().st_size)
        for path in paths
    )


def load_records() -> list[RagRecord]:
    """Reload only when a file changes; otherwise use the memory cache."""
    global _cached_signature
    global _cached_records

    paths = _source_paths()
    signature = _signature(paths)

    with _cache_lock:
        if signature == _cached_signature:
            return list(_cached_records)

        records: list[RagRecord] = []

        for path in paths:
            suffix = path.suffix.lower()

            try:
                if suffix == ".csv":
                    records.extend(_load_csv(path))
                elif suffix == ".xlsx":
                    records.extend(_load_xlsx(path))
                elif suffix == ".txt":
                    records.extend(_load_txt(path))
                else:
                    print(f"[RAG WARNING] Unsupported format: {path.name}")
            except Exception as error:
                print(f"[RAG LOAD ERROR] {path}: {error}")

        _cached_signature = signature
        _cached_records = records

        print(f"[RAG] Loaded {len(records)} rows")
        return list(records)


def _score(query: str, query_tokens: set[str], record: RagRecord) -> float:
    record_tokens = _tokens(record.text)

    if not record_tokens:
        return 0.0

    overlap = len(query_tokens & record_tokens)
    semantic_score = overlap / math.sqrt(
        max(1, len(query_tokens)) * max(1, len(record_tokens))
    )

    score = semantic_score
    query_lower = _normalise(query)

    module_id = _normalise(record.fields.get("module_id", ""))
    function_name = _normalise(record.fields.get("function", ""))
    keywords = _tokens(record.fields.get("keywords", ""))

    if module_id and module_id in query_lower:
        score += 4.0

    if function_name and function_name in query_lower:
        score += 2.0

    if keywords:
        score += 1.5 * len(query_tokens & keywords) / max(1, len(keywords))

    if record.always_include:
        score += 100.0

    return score


def retrieve(query: str) -> list[tuple[float, RagRecord]]:
    if not config.RAG_ENABLED:
        return []

    records = load_records()
    query_tokens = _tokens(query)

    if not records or not query_tokens:
        return []

    scored = [(_score(query, query_tokens, record), record) for record in records]
    scored.sort(key=lambda item: item[0], reverse=True)

    selected: list[tuple[float, RagRecord]] = []

    for score, record in scored:
        if record.always_include or score >= config.RAG_MIN_SCORE:
            selected.append((score, record))

        if len(selected) >= config.RAG_TOP_K:
            break

    return selected


def build_rag_context(query: str) -> str:
    selected = retrieve(query)

    if not selected:
        return ""

    blocks: list[str] = []
    current_size = 0

    for index, (score, record) in enumerate(selected, start=1):
        block = (
            f"[RAG Data {index}]\n"
            f"Source: {record.source} / {record.location}\n"
            f"Content: {record.text}"
        )

        if current_size + len(block) > config.RAG_MAX_CONTEXT_CHARS:
            break

        blocks.append(block)
        current_size += len(block)

        print(
            f"[RAG MATCH {index}] score={score:.3f} "
            f"{record.source}/{record.location}"
        )

    return "\n\n".join(blocks)
