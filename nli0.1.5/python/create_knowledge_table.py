from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


PROJECT_ROOT = Path(__file__).resolve().parent.parent
OUTPUT_PATH = PROJECT_ROOT / "knowledge" / "module_knowledge.xlsx"

HEADERS = [
    "id",
    "category",
    "module_id",
    "module_name",
    "function",
    "parameters",
    "description",
    "conditions",
    "safety",
    "keywords",
    "example",
    "always_include",
    "enabled",
]

SAMPLE_ROWS = [
    [
        "GLOBAL-001", "global_rule", "none", "Common safety rules", "none", "none",
        "Use only registered modules and functions; do not create commands that are not in the data.",
        "All requests", "If dangerous or ambiguous, prioritize stopping or asking a confirmation question over execution.",
        "common safety danger ambiguous confirm stop", "pick that up → ask again which object", True, True,
    ],
    [
        "ARM-001", "module_function", "ARM01", "Robot arm", "move_to",
        "target", "Move the arm end effector to the specified target or position.",
        "When the target is clear and the movement path is safe", "Do not execute if there is a possibility of collision.",
        "arm move pick cup object position", "move_to(target=cup)", False, True,
    ],
    [
        "ARM-002", "module_function", "ARM01", "Robot arm", "grip",
        "state=open|close", "Open or close the gripper.",
        "When the target is within gripper range", "Do not apply excessive force.",
        "gripper grab release open close", "grip(state=close)", False, True,
    ],
    [
        "ARM-003", "module_function", "ARM01", "Robot arm", "stop",
        "none", "Immediately stop arm movement.",
        "Danger, collision possibility, user cancellation", "May be used with priority as a safe alternative command.",
        "arm stop halt danger collision", "stop()", False, True,
    ],
    [
        "WHEEL-001", "module_function", "WHEEL01", "Drive wheels", "forward",
        "speed", "Move the robot forward.",
        "When there is no obstacle ahead and position and battery are normal", "Do not execute if the battery is low or the position is unknown.",
        "forward ahead move wheels", "forward(speed=30)", False, True,
    ],
    [
        "WHEEL-002", "module_function", "WHEEL01", "Drive wheels", "stop",
        "none", "Immediately stop the wheels.",
        "Danger or user stop request", "Choose with priority when movement is uncertain.",
        "stop halt wheels danger", "stop()", False, True,
    ],
]


def main() -> None:
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Knowledge"

    worksheet.append(HEADERS)
    for row in SAMPLE_ROWS:
        worksheet.append(row)

    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions

    widths = {
        "A": 16, "B": 18, "C": 14, "D": 18, "E": 18, "F": 22,
        "G": 42, "H": 42, "I": 42, "J": 30, "K": 32, "L": 15, "M": 12,
    }
    for column, width in widths.items():
        worksheet.column_dimensions[column].width = width

    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    table = Table(displayName="KnowledgeTable", ref=worksheet.dimensions)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    worksheet.add_table(table)

    workbook.save(OUTPUT_PATH)
    print(f"Knowledge table creation complete: {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
